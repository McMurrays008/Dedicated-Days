from __future__ import annotations
import csv, gc, json, os, re, shutil
from datetime import date, datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

HERE = Path(__file__).resolve().parent
DOWNLOAD_DIR = HERE / "downloads"
DATA_FILE = HERE / "dedicated-day-data.json"
FAILURE_SCREENSHOT = HERE / "tpn_failure.png"
FAILURE_HTML = HERE / "tpn_failure.html"
DOWNLOAD_DIR.mkdir(exist_ok=True)

RISK_STATUSES = {"ACH","ATH","ATHP","ICC","ITD","ITDP","ITH","ITHP","OFDP","WCD","WCDP","WDDP"}

def stage(cb, name):
    print(f"[collector] {name}", flush=True)
    if cb: cb(name)

def norm(v):
    return re.sub(r"\s+"," ",str(v or "").strip())

def norm_docket(v):
    s=norm(v)
    if s.endswith(".0"): s=s[:-2]
    return re.sub(r"\s+","",s).upper()

def depot_code(v):
    s=norm(v)
    if not s: return ""
    try: return str(int(float(s)))
    except Exception: return s.lstrip("0") or "0"

def env_holidays():
    out=set()
    for raw in os.getenv("TPN_HOLIDAYS","").split(","):
        raw=raw.strip()
        if raw:
            try: out.add(datetime.strptime(raw,"%Y-%m-%d").date())
            except ValueError: pass
    return out

def most_recent_working_day(today, holidays):
    # Preserve the established "last working day" behaviour by default.
    if os.getenv("DATE_MODE","last_working_day").lower()=="today":
        d=today
    else:
        d=today-timedelta(days=1)
    while d.weekday()>=5 or d in holidays:
        d-=timedelta(days=1)
    return d

def working_range(end, count, holidays):
    days=[]; d=end
    while len(days)<count:
        if d.weekday()<5 and d not in holidays: days.append(d)
        d-=timedelta(days=1)
    return min(days),max(days)

def dtxt(d): return d.strftime(os.getenv("DATE_FORMAT","%d/%m/%Y"))

def first_visible(page, selectors):
    for sel in selectors:
        try:
            loc=page.locator(sel)
            if loc.count() and loc.first.is_visible(): return loc.first
        except Exception: pass
    return None

def click_named(page, name, exact=True):
    candidates=[
        page.get_by_role("menuitem",name=name,exact=exact),
        page.get_by_role("link",name=name,exact=exact),
        page.get_by_role("button",name=name,exact=exact),
        page.get_by_text(name,exact=exact),
    ]
    for loc in candidates:
        try:
            if loc.count() and loc.first.is_visible():
                loc.first.click(timeout=12000)
                return
        except Exception: pass
    raise RuntimeError(f"Could not find visible control: {name}")

def click_top_nav(page, name):
    """Open a Pilot TPN top-nav item robustly.

    Pilot can render the navigation differently between sessions. Try the
    normal visible controls first, then Browse-specific menu text/onclick
    elements, then navigate back to /Dashboard and retry once.
    """
    rx=re.compile(rf"^\s*{re.escape(name)}\s*$",re.I)

    def contexts():
        return [page]+[f for f in page.frames if f is not page.main_frame]

    def try_click():
        selectors=[
            f"a:has-text('{name}')",
            f"button:has-text('{name}')",
            f"[role='menuitem']:has-text('{name}')",
            f"[role='link']:has-text('{name}')",
            f"[onclick*='{name}' i]",
        ]
        for ctx in contexts():
            # Exact accessible text first.
            for getter in (
                lambda: ctx.get_by_role("link",name=rx),
                lambda: ctx.get_by_role("button",name=rx),
                lambda: ctx.get_by_role("menuitem",name=rx),
                lambda: ctx.get_by_text(rx),
            ):
                try:
                    q=getter()
                    for i in range(min(q.count(),20)):
                        el=q.nth(i)
                        if el.is_visible():
                            try: el.click(force=True,timeout=5000)
                            except Exception: el.evaluate("(el)=>el.click()")
                            return True
                except Exception:
                    pass

            # Forgiving CSS candidates.
            for sel in selectors:
                try:
                    q=ctx.locator(sel)
                    for i in range(min(q.count(),30)):
                        el=q.nth(i)
                        if not el.is_visible(): continue
                        txt=(el.inner_text() or "").strip()
                        if name.lower() not in txt.lower() and name.lower() not in str(el.get_attribute("onclick") or "").lower():
                            continue
                        try: el.click(force=True,timeout=5000)
                        except Exception: el.evaluate("(el)=>el.click()")
                        return True
                except Exception:
                    pass

        # Browse has historically exposed tabBrowseChange() even when the
        # top menu text itself is not discoverable.
        if name.lower()=="browse":
            for ctx in contexts():
                try:
                    q=ctx.locator("[onclick*='tabBrowseChange' i]")
                    for i in range(min(q.count(),20)):
                        el=q.nth(i)
                        if el.is_visible():
                            el.evaluate("(el)=>el.click()")
                            return True
                except Exception:
                    pass
        return False

    if try_click():
        page.wait_for_timeout(900)
        return

    # Some TPN sessions land on a child/blank page after authentication.
    # Return to the known Dashboard URL and try the navigation again.
    try:
        base=os.getenv("TPN_LOGIN_URL","https://pilot.tpnconnect.com/").rstrip("/")
        page.goto(base+"/Dashboard",wait_until="domcontentloaded",timeout=30000)
        page.wait_for_timeout(1200)
        if try_click():
            page.wait_for_timeout(900)
            return
    except Exception:
        pass

    # Put actionable diagnostics directly into the Render log.
    visible=[]
    for ctx in contexts():
        for sel in ("a","button","[role='menuitem']","[role='link']","[onclick]"):
            try:
                q=ctx.locator(sel)
                for i in range(min(q.count(),80)):
                    el=q.nth(i)
                    if not el.is_visible(): continue
                    txt=" ".join((el.inner_text() or "").split())
                    onclick=el.get_attribute("onclick")
                    if txt or onclick:
                        visible.append({"tag":sel,"text":txt[:100],"onclick":(onclick or "")[:120]})
            except Exception:
                pass
    try: current_url=page.url
    except Exception: current_url="?"
    try: title=page.title()
    except Exception: title="?"
    print(f"[collector] NAV DIAGNOSTIC url={current_url!r} title={title!r} visible={visible[:120]}",flush=True)
    raise RuntimeError(f"Could not find visible top navigation control: {name}")

def fill_date(page, which, value):
    """Find TPN date fields across the main page and child frames.

    Pilot TPN Browse does not consistently use the DateFrom/DateTo names used
    on Integration, so this also recognises labels, ids/names containing
    from/to + date, and the first input following a visible Date From/To label.
    """
    want_from="from" in which.lower()
    stem="DateFrom" if want_from else "DateTo"
    word="from" if want_from else "to"
    contexts=[page]+[f for f in page.frames if f is not page.main_frame]

    selectors=[
        f"input[name='{stem}']",
        f"input[name*='{stem}' i]",
        f"input[id*='{stem}' i]",
        f"input[name*='date'][name*='{word}' i]",
        f"input[id*='date'][id*='{word}' i]",
        f"input[name*='{word}' i][name*='date' i]",
        f"input[id*='{word}' i][id*='date' i]",
        f"input[placeholder*='{which}' i]",
        f"input[aria-label*='{which}' i]",
    ]

    loc=None
    for ctx in contexts:
        loc=first_visible(ctx,selectors)
        if loc: break

        # Accessible labels, where present.
        try:
            q=ctx.get_by_label(re.compile(rf"Date\s*{word}",re.I))
            for i in range(q.count()):
                if q.nth(i).is_visible():
                    loc=q.nth(i); break
        except Exception:
            pass
        if loc: break

        # TPN/Kendo pages sometimes render a plain text label with an
        # unlabelled input immediately after it.
        try:
            labels=ctx.get_by_text(re.compile(rf"^\s*Date\s*{word}\s*:?\s*$",re.I))
            for i in range(min(labels.count(),10)):
                lab=labels.nth(i)
                if not lab.is_visible(): continue
                q=lab.locator("xpath=following::input[not(@type='hidden')][1]")
                if q.count() and q.first.is_visible():
                    loc=q.first; break
        except Exception:
            pass
        if loc: break

    if not loc:
        # Produce useful diagnostics in Render logs if TPN changes again.
        found=[]
        for ctx in contexts:
            try:
                for i in range(min(ctx.locator("input").count(),40)):
                    el=ctx.locator("input").nth(i)
                    if not el.is_visible(): continue
                    found.append({
                        "name":el.get_attribute("name"),
                        "id":el.get_attribute("id"),
                        "type":el.get_attribute("type"),
                        "placeholder":el.get_attribute("placeholder"),
                        "aria":el.get_attribute("aria-label"),
                    })
            except Exception:
                pass
        print(f"[collector] Visible inputs while looking for {which}: {found}",flush=True)
        raise RuntimeError(f"Could not identify {which} field")

    # Kendo/date-picker inputs can be readonly or have JS handlers. Fill first;
    # fall back to DOM value assignment plus input/change events.
    try:
        loc.fill(value,force=True)
    except Exception:
        loc.evaluate("""(el, value) => {
            el.removeAttribute('readonly');
            el.value=value;
            el.dispatchEvent(new Event('input',{bubbles:true}));
            el.dispatchEvent(new Event('change',{bubbles:true}));
        }""",value)
    try: loc.press("Tab")
    except Exception: pass

def login(page, cb):
    username=os.getenv("TPN_USERNAME")
    password=os.getenv("TPN_PASSWORD")
    if not username or not password:
        raise RuntimeError("TPN_USERNAME and TPN_PASSWORD must be set in Render environment variables.")
    stage(cb,"Opening TPN login")
    page.goto(os.getenv("TPN_LOGIN_URL","https://connect.tpnsecure.com/"),wait_until="domcontentloaded",timeout=45000)
    u=first_visible(page,["input[placeholder*='Username' i]","input[name='Username']","input[name='username']","#Username","#username","input[type='text']"])
    pw=first_visible(page,["input[placeholder*='Password' i]","input[name='Password']","input[name='password']","#Password","#password","input[type='password']"])
    # Already signed in is also acceptable.
    if not u and not pw:
        if page.get_by_text("Browse",exact=True).count(): return
        raise RuntimeError("Could not identify TPN login controls.")
    stage(cb,"Submitting TPN login")
    u.fill(username); pw.fill(password)
    try: click_named(page,"Login",True)
    except Exception: pw.press("Enter")
    page.wait_for_timeout(1600)
    if page.locator("input[type='password']:visible").count():
        raise RuntimeError("TPN login page remained visible after submitting credentials.")

def click_export(page, regex, target, cb):
    stage(cb,f"Downloading {target.name}")
    loc=None
    for candidate in [
        page.get_by_role("button",name=re.compile(regex,re.I)),
        page.get_by_role("link",name=re.compile(regex,re.I)),
        page.get_by_text(re.compile(regex,re.I))
    ]:
        try:
            if candidate.count() and candidate.first.is_visible():
                loc=candidate.first; break
        except Exception: pass
    if not loc: raise RuntimeError(f"Could not find export control matching {regex}")
    with page.expect_download(timeout=120000) as di:
        try: loc.click(timeout=15000,no_wait_after=True,force=True)
        except Exception: loc.evaluate("el=>el.click()")
    di.value.save_as(target)
    return target

def export_integration(page,start,end,cb):
    stage(cb,"Opening Integration")

    # The Pilot site keeps the Integration page under /Dashboard and opens it
    # as an in-page tab.  Navigating to Dashboard first makes the top menu
    # predictable after login, then the forgiving nav selector handles the
    # icon+text control used by TPN Connect.
    base=os.getenv("TPN_LOGIN_URL","https://pilot.tpnconnect.com/").rstrip("/")
    dashboard_url=os.getenv("TPN_DASHBOARD_URL",base+"/Dashboard")
    if "/dashboard" not in page.url.lower():
        try:
            page.goto(dashboard_url,wait_until="domcontentloaded",timeout=45000)
            page.wait_for_timeout(1200)
        except Exception:
            pass
    click_top_nav(page,"Integration")
    page.wait_for_timeout(1200)

    stage(cb,"Setting Integration date range")
    fill_date(page,"Date From",dtxt(start)); fill_date(page,"Date To",dtxt(end))

    # Requestor is the required export type on the Pilot Integration screen.
    try:
        radio=page.get_by_label(re.compile(r"^Requestor$",re.I))
        if radio.count() and radio.first.is_visible():
            radio.first.check(force=True)
    except Exception:
        try:
            radio=page.get_by_text(re.compile(r"^Requestor$",re.I))
            if radio.count() and radio.first.is_visible(): radio.first.click(force=True)
        except Exception:
            pass

    target=DOWNLOAD_DIR/f"RAW_ConsignmentExport_{datetime.now():%Y%m%d_%H%M%S}.csv"
    return click_export(page,r"^Export$",target,cb)

def export_browse(page,start,end,cb):
    stage(cb,"Opening Browse")

    # Pilot TPN Connect uses Browse as a dropdown.  Clicking the top-level
    # Browse control only opens that dropdown; the actual browse screen used
    # for the export is the "Browse Discreps" submenu item.  Use forced/DOM
    # clicks because the dropdown itself can otherwise intercept pointer events.
    click_top_nav(page,"Browse")
    page.wait_for_timeout(500)
    submenu=None
    pattern=re.compile(r"Browse\s+Discrep",re.I)
    for candidate in [
        page.get_by_role("link",name=pattern),
        page.get_by_role("menuitem",name=pattern),
        page.locator("a").filter(has_text=pattern),
        page.locator("[onclick]").filter(has_text=pattern),
        page.get_by_text(pattern),
    ]:
        try:
            for i in range(min(candidate.count(),8)):
                item=candidate.nth(i)
                if item.is_visible():
                    submenu=item
                    break
            if submenu: break
        except Exception:
            pass
    if not submenu:
        raise RuntimeError("Could not find visible Browse Discreps submenu item")
    try: submenu.click(timeout=12000,force=True)
    except Exception: submenu.evaluate("el=>el.click()")
    page.wait_for_timeout(1000)

    stage(cb,"Setting Browse date range")
    fill_date(page,"Date From",dtxt(start)); fill_date(page,"Date To",dtxt(end))
    stage(cb,"Loading Browse results")

    # Pilot TPN has used different labels/types for the content-area submit
    # control. Search all frames, excluding the top-nav tabBrowseChange control.
    search=None
    contexts=[page]+[f for f in page.frames if f is not page.main_frame]
    wanted=re.compile(r"^(browse|search|view|go|submit|load|find|refresh)$",re.I)
    # Current Pilot Browse uses an unlabeled button whose action is fuzzySearch().
    for ctx in contexts:
        try:
            q=ctx.locator("[onclick*='fuzzySearch' i]")
            for i in range(min(q.count(),20)):
                item=q.nth(i)
                if item.is_visible():
                    search=item
                    break
        except Exception:
            pass
        if search:
            break

    for ctx in contexts:
        if search:
            break
        for sel in ("input[type='submit']","button[type='submit']","input[type='button']",
                    "input","button","[onclick]"):
            try:
                q=ctx.locator(sel)
                for i in range(min(q.count(),100)):
                    item=q.nth(i)
                    if not item.is_visible(): continue
                    onclick=(item.get_attribute("onclick") or "").lower()
                    if "tabbrowsechange" in onclick: continue
                    typ=(item.get_attribute("type") or "").lower()
                    value=(item.get_attribute("value") or "").strip()
                    txt=" ".join((item.inner_text() or "").split()).strip()
                    aria=(item.get_attribute("aria-label") or "").strip()
                    label=value or txt or aria
                    if typ=="submit" or wanted.match(label):
                        search=item; break
            except Exception:
                pass
            if search: break
        if search: break

    if not search:
        # Last resort: submit the form containing the Browse date controls.
        for ctx in contexts:
            try:
                forms=ctx.locator("form")
                for i in range(forms.count()):
                    form=forms.nth(i)
                    body=(form.inner_text() or "").lower()
                    if "date from" in body and "date to" in body:
                        form.evaluate("(f)=>f.requestSubmit ? f.requestSubmit() : f.submit()")
                        page.wait_for_timeout(1400)
                        search=True
                        break
            except Exception:
                pass
            if search: break

    if not search:
        found=[]
        for ctx in contexts:
            for sel in ("button","input","[onclick]"):
                try:
                    q=ctx.locator(sel)
                    for i in range(min(q.count(),100)):
                        el=q.nth(i)
                        if not el.is_visible(): continue
                        found.append({"tag":sel,"type":el.get_attribute("type"),
                                      "value":el.get_attribute("value"),
                                      "text":" ".join((el.inner_text() or "").split())[:80],
                                      "onclick":(el.get_attribute("onclick") or "")[:120]})
                except Exception: pass
        print(f"[collector] BROWSE ACTION DIAGNOSTIC visible={found[:150]}",flush=True)
        raise RuntimeError("Could not find Browse results/search button")

    if search is not True:
        try: search.click(timeout=15000,force=True)
        except Exception: search.evaluate("el=>el.click()")
        page.wait_for_timeout(1400)
    target=DOWNLOAD_DIR/f"RAW_TPN_Dedicated_Day_Browse_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    stage(cb,f"Downloading {target.name}")

    # Browse's "Export To Excel" text can sit inside a non-clickable wrapper.
    # Find the actual actionable element across the page/frames and prefer
    # href/onclick/button controls over a plain text container.
    export=None
    export_context=None
    export_rx=re.compile(r"Export\s*To\s*Excel",re.I)
    for ctx in contexts:
        for sel in ("a","button","input[type='button']","input[type='submit']","[onclick]"):
            try:
                q=ctx.locator(sel)
                for i in range(min(q.count(),120)):
                    el=q.nth(i)
                    if not el.is_visible(): continue
                    label=" ".join(((el.inner_text() or "")+" "+(el.get_attribute("value") or "")+" "+(el.get_attribute("aria-label") or "")).split())
                    if export_rx.search(label):
                        export=el; export_context=ctx; break
            except Exception:
                pass
            if export: break
        if export: break

    # If only the text node is discoverable, walk up to its nearest clickable
    # ancestor instead of clicking the wrapper itself.
    if not export:
        for ctx in contexts:
            try:
                q=ctx.get_by_text(export_rx)
                for i in range(min(q.count(),20)):
                    el=q.nth(i)
                    if not el.is_visible(): continue
                    clickable=el.locator("xpath=ancestor-or-self::*[self::a or self::button or @onclick][1]")
                    if clickable.count() and clickable.first.is_visible():
                        export=clickable.first; export_context=ctx; break
            except Exception:
                pass
            if export: break

    if not export:
        found=[]
        for ctx in contexts:
            for sel in ("a","button","input","[onclick]"):
                try:
                    q=ctx.locator(sel)
                    for i in range(min(q.count(),120)):
                        el=q.nth(i)
                        if not el.is_visible(): continue
                        found.append({"tag":sel,
                                      "text":" ".join((el.inner_text() or "").split())[:100],
                                      "value":el.get_attribute("value"),
                                      "href":el.get_attribute("href"),
                                      "onclick":(el.get_attribute("onclick") or "")[:160]})
                except Exception: pass
        print(f"[collector] BROWSE EXPORT DIAGNOSTIC visible={found[:180]}",flush=True)
        raise RuntimeError("Could not find actionable Browse Export To Excel control")

    print(f"[collector] Browse export control href={export.get_attribute('href')!r} onclick={export.get_attribute('onclick')!r}",flush=True)
    try:
        with page.expect_download(timeout=120000) as di:
            try: export.click(timeout=15000,force=True,no_wait_after=True)
            except Exception: export.evaluate("el=>el.click()")
        di.value.save_as(target)
        return target
    except Exception:
        # Log the exact export-area controls if TPN did not emit a browser
        # download event, so the next correction is based on the real action.
        found=[]
        for ctx in contexts:
            for sel in ("a","button","input","[onclick]"):
                try:
                    q=ctx.locator(sel)
                    for i in range(min(q.count(),120)):
                        el=q.nth(i)
                        if not el.is_visible(): continue
                        label=" ".join(((el.inner_text() or "")+" "+(el.get_attribute("value") or "")).split())
                        if "export" in label.lower() or "excel" in label.lower() or "export" in (el.get_attribute("onclick") or "").lower():
                            found.append({"tag":sel,"text":label[:120],
                                          "href":el.get_attribute("href"),
                                          "onclick":(el.get_attribute("onclick") or "")[:200]})
                except Exception: pass
        print(f"[collector] BROWSE EXPORT DIAGNOSTIC controls={found[:100]}",flush=True)
        raise

def read_csv(path):
    raw=Path(path).read_bytes()
    for enc in ("utf-8-sig","cp1252","latin1"):
        try:
            text=raw.decode(enc); break
        except UnicodeDecodeError: continue
    return list(csv.DictReader(text.splitlines()))

def read_xlsx(path):
    wb=load_workbook(path,read_only=True,data_only=True)
    ws=wb.active
    rows=ws.iter_rows(values_only=True)
    headers=[norm(x) for x in next(rows)]
    out=[]
    for vals in rows:
        if any(v is not None and str(v).strip() for v in vals):
            out.append(dict(zip(headers,vals)))
    wb.close()
    return out

def first(r,*names):
    lower={norm(k).lower():v for k,v in r.items()}
    for n in names:
        key=norm(n).lower()
        if key in lower and norm(lower[key])!="": return lower[key]
    return ""

def num(v):
    try:return int(float(str(v).strip()))
    except:return 0

def pallet_count(r):
    total=0
    for k,v in r.items():
        lk=norm(k).lower()
        if "pallet" in lk and "type" not in lk and "weight" not in lk and lk not in {"pallets","pallet"}:
            try: total+=int(float(str(v).strip() or 0))
            except: pass
    if total: return total
    return num(first(r,"Pallets","Pallet"))

def filter_browse(rows):
    prefix=os.getenv("SERVICE_PREFIX","DD").upper()
    kept=[]
    for r in rows:
        service=norm(first(r,"Service")).upper()
        req=first(r,"Req","Request","Requesting Depot")
        dele=first(r,"Del","Deliver","Delivery Depot")
        if not service.startswith(prefix): continue
        if req and depot_code(req)!="8": continue
        # 10 September method: Delivery Depot 8 is deliberately INCLUDED.
        kept.append(r)
    return kept

def filter_integration(rows):
    prefix=os.getenv("SERVICE_PREFIX","DD").upper()
    return [r for r in rows if norm(first(r,"Service")).upper().startswith(prefix)]

def parse_delivery_date(v):
    if isinstance(v,datetime): return v.date()
    if isinstance(v,date): return v
    if isinstance(v,(int,float)):
        try:
            from openpyxl.utils.datetime import from_excel
            return from_excel(v).date()
        except Exception:return None
    s=norm(v)
    for f in ("%d/%m/%Y","%Y-%m-%d","%d-%m-%Y","%d/%m/%y"):
        try:return datetime.strptime(s,f).date()
        except ValueError:pass
    return None

def status_lookup_from_integration(rows):
    m={}
    for r in rows:
        d=norm_docket(first(r,"Docket","Consignment","Consignment Number"))
        if d: m[d]=norm(first(r,"Status","STATUS","Status Code"))
    return m

def make_dashboard_rows(browse_rows,status_lookup=None,target_date=None):
    out=[]
    for r in browse_rows:
        ddate=parse_delivery_date(first(r,"Delivery Date"))
        if target_date and ddate and ddate!=target_date: continue
        docket=norm_docket(first(r,"Docket","Consignment","Consignment Number"))
        if not docket: continue
        status=(status_lookup or {}).get(docket) or norm(first(r,"STATUS","Status","Status Code"))
        out.append({
            "Docket":docket,
            "Sender":norm(first(r,"Consignor Name","Sender","Consignor")),
            "Service":norm(first(r,"Service")),
            "Status":status,
            "Pallets":pallet_count(r),
            "Consignee":norm(first(r,"Consignee Name","Consignee")),
            "Postcode":norm(first(r,"Consignee Postcode","Post Code","Postcode")),
            "Delivery Depot":norm(first(r,"Delivery Depot","Deliver","Del")),
        })
    return out

def atomic_write(payload):
    tmp=DATA_FILE.with_suffix(".json.tmp")
    tmp.write_text(json.dumps(payload,ensure_ascii=False,indent=2,default=str),encoding="utf-8")
    tmp.replace(DATA_FILE)


def manual_morning_import(xlsx_path, stage_callback=None):
    """Initialise today's fixed delivery population from the manually supplied
    TPN Dedicated Day Check. Statuses are whatever is present in that file;
    later Browse refreshes update Status only and never add/remove deliveries.
    """
    stage(stage_callback,"Reading manual Dedicated Day Check")
    browse=filter_browse(read_xlsx(xlsx_path))
    # The manually uploaded morning file may contain future delivery dates.
    # The fixed population must be for the day the file is imported, not the
    # maximum Delivery Date present in the workbook.
    target=datetime.now(ZoneInfo(os.getenv("TIMEZONE","Europe/London"))).date()
    rows=make_dashboard_rows(browse,status_lookup=None,target_date=target)
    if not rows:
        raise RuntimeError("No Dedicated Day rows survived the 10 September collection rules.")
    payload={
        "delivery_date":target.isoformat(),
        "generated_at":datetime.now(ZoneInfo(os.getenv("TIMEZONE","Europe/London"))).isoformat(timespec="seconds"),
        "mode":"manual_morning_import",
        "source":{
            "manual_browse":Path(xlsx_path).name,
            "browse_dd_rows":len(browse),
            "fixed_population":len(rows)
        },
        "rows":rows
    }
    atomic_write(payload)
    stage(stage_callback,f"Morning import complete: {len(rows)} deliveries fixed for {target.isoformat()}")
    return {"delivery_date":target.isoformat(),"rows":len(rows),"pallets":sum(num(r.get("Pallets")) for r in rows)}

def full_refresh(stage_callback=None):
    holidays=env_holidays()
    end=most_recent_working_day(date.today(),holidays)
    start,end=working_range(end,int(os.getenv("WORKING_DAYS_TO_EXPORT","7")),holidays)
    stage(stage_callback,f"Using date range {dtxt(start)} to {dtxt(end)}")
    integration_path=browse_path=None
    with sync_playwright() as p:
        browser=p.chromium.launch(headless=os.getenv("HEADLESS","true").lower()=="true",
            args=["--no-sandbox","--disable-setuid-sandbox","--disable-dev-shm-usage","--disable-gpu","--disable-extensions","--renderer-process-limit=1"])
        context=browser.new_context(accept_downloads=True,viewport={"width":1100,"height":760})
        page=context.new_page()
        page.route("**/*",lambda route: route.abort() if route.request.resource_type in {"image","media","font"} else route.continue_())
        page.set_default_timeout(15000)
        try:
            login(page,stage_callback)
            integration_path=export_integration(page,start,end,stage_callback)
            browse_path=export_browse(page,start,end,stage_callback)
        except Exception:
            try: page.screenshot(path=str(FAILURE_SCREENSHOT),full_page=False)
            except Exception: pass
            try: FAILURE_HTML.write_text(page.content(),encoding="utf-8")
            except Exception: pass
            raise
        finally:
            context.close(); browser.close(); gc.collect()

    stage(stage_callback,"Filtering exports locally")
    integration=filter_integration(read_csv(integration_path))
    browse=filter_browse(read_xlsx(browse_path))
    dates=[parse_delivery_date(first(r,"Delivery Date")) for r in browse]
    dates=[d for d in dates if d]
    target=max(dates) if dates else end
    statuses=status_lookup_from_integration(integration)
    rows=make_dashboard_rows(browse,statuses,target)
    if browse and not rows:
        raise RuntimeError("Browse export had rows but zero dashboard rows survived the delivery-date filter.")
    matched=sum(1 for r in rows if r["Docket"] in statuses)
    payload={
        "delivery_date":target.isoformat(),
        "generated_at":datetime.now(ZoneInfo(os.getenv("TIMEZONE","Europe/London"))).isoformat(timespec="seconds"),
        "mode":"full",
        "source":{
            "integration":Path(integration_path).name,
            "browse":Path(browse_path).name,
            "range_start":start.isoformat(),"range_end":end.isoformat(),
            "integration_dd_rows":len(integration),"browse_dd_rows":len(browse),
            "status_matches":matched
        },
        "rows":rows
    }
    atomic_write(payload)
    stage(stage_callback,f"Completed full refresh: {len(rows)} deliveries; {matched} statuses matched")
    return payload["source"] | {"delivery_date":payload["delivery_date"],"rows":len(rows)}

def status_refresh(stage_callback=None):
    if not DATA_FILE.exists():
        raise RuntimeError("No morning Dedicated Day Check has been imported yet.")
    current=json.loads(DATA_FILE.read_text(encoding="utf-8"))
    if not current.get("rows"):
        raise RuntimeError("The morning import contains no deliveries.")

    try:
        target=datetime.strptime(current["delivery_date"],"%Y-%m-%d").date()
    except Exception:
        raise RuntimeError("The current dashboard has no valid delivery_date.")

    # Search Browse for the fixed delivery day only. This refresh is status-only:
    # it cannot add or remove deliveries from the morning population.
    stage(stage_callback,f"Refreshing Browse statuses for {dtxt(target)}")
    with sync_playwright() as p:
        browser=p.chromium.launch(headless=os.getenv("HEADLESS","true").lower()=="true",
            args=["--no-sandbox","--disable-setuid-sandbox","--disable-dev-shm-usage","--disable-gpu","--disable-extensions","--renderer-process-limit=1"])
        context=browser.new_context(accept_downloads=True,viewport={"width":1100,"height":760})
        page=context.new_page()
        page.route("**/*",lambda route: route.abort() if route.request.resource_type in {"image","media","font"} else route.continue_())
        page.set_default_timeout(15000)
        try:
            login(page,stage_callback)
            browse_path=export_browse(page,target,target,stage_callback)
        except Exception:
            try: page.screenshot(path=str(FAILURE_SCREENSHOT),full_page=False)
            except Exception: pass
            try: FAILURE_HTML.write_text(page.content(),encoding="utf-8")
            except Exception: pass
            raise
        finally:
            context.close(); browser.close(); gc.collect()

    stage(stage_callback,"Matching Browse statuses to morning Dockets")
    browse=filter_browse(read_xlsx(browse_path))
    latest={}
    for r in browse:
        docket=norm_docket(first(r,"Docket","Consignment","Consignment Number"))
        if docket:
            latest[docket]=norm(first(r,"STATUS","Status","Status Code"))

    changed=matched=0
    for r in current["rows"]:
        d=norm_docket(r.get("Docket"))
        if d in latest and latest[d]:
            matched+=1
            if latest[d] != norm(r.get("Status")):
                r["Status"]=latest[d]
                changed+=1

    current["generated_at"]=datetime.now(ZoneInfo(os.getenv("TIMEZONE","Europe/London"))).isoformat(timespec="seconds")
    current["mode"]="status"
    current["status_refresh"]={
        "matched_dockets":matched,
        "changed_statuses":changed,
        "source":Path(browse_path).name,
        "delivery_date":target.isoformat()
    }
    atomic_write(current)
    stage(stage_callback,f"Completed status refresh: {matched} matched; {changed} changed")
    return current["status_refresh"]

